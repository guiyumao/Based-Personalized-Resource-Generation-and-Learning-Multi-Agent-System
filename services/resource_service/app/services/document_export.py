"""Create a self-contained Word document from stored learning content."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from services.resource_service.app.schemas.resource import ResourceItem

DOCUMENT_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PROSE_TYPES = {"courseware", "notes", "exam"}
TEXT_FORMATS = {"markdown", "text", "txt"}


def _text(value: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "\ufffd", value)


def _inline(paragraph, text: str) -> None:
    for part in re.split(r"(`[^`\n]+`|\*\*[^*\n]+\*\*)", _text(text)):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(part)


def build_resource_docx(resource: ResourceItem) -> bytes:
    """Render headings, lists, tables and code; preserve formulas as LaTeX text.

    No embedded links, code or remote images are executed or downloaded.
    """
    if resource.source_type != "generated" or resource.type not in PROSE_TYPES or resource.format not in TEXT_FORMATS:
        raise ValueError("此资源类型暂不支持 Word 导出。")
    content = (resource.content or "").strip()
    if not content:
        raise ValueError("该资源没有可导出的正文。")
    if len(content) > 500_000:
        raise ValueError("资源正文过长，暂不支持 Word 导出。")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(_text(resource.title), 0)
    document.add_paragraph(_text(f"知识点：{resource.knowledge_point}　资源编号：{resource.id}"))
    document.add_paragraph("公式按 LaTeX 原文保留；代码仅作文本展示。")
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith(("```", "~~~")):
            fence = stripped[:3]
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(fence):
                code.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run(_text("\n".join(code)))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|") and lines[index].strip().endswith("|"):
                cells = [cell.strip() for cell in lines[index].strip()[1:-1].split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                table = document.add_table(rows=0, cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for row in rows:
                    cells = table.add_row().cells
                    for column, value in enumerate(row):
                        _inline(cells[column].paragraphs[0], value)
            continue
        elif heading := re.match(r"^(#{1,6})\s+(.+)$", stripped):
            paragraph = document.add_heading(level=min(len(heading[1]), 3))
            _inline(paragraph, heading[2])
        elif bullet := re.match(r"^[-*+]\s+(.+)$", stripped):
            _inline(document.add_paragraph(style="List Bullet"), bullet[1])
        elif numbered := re.match(r"^\d+[.)]\s+(.+)$", stripped):
            _inline(document.add_paragraph(style="List Number"), numbered[1])
        elif stripped:
            _inline(document.add_paragraph(), line)
        index += 1
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
