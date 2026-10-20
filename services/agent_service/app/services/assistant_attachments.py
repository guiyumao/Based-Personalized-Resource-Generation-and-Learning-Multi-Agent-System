"""Bounded text extraction for documents explicitly attached to a question."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import PurePath

from fastapi import HTTPException
from sqlalchemy.orm import Session

from common.models.learning import AssistantAttachment

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 20_000
MAX_PROMPT_CHARS = 12_000
MAX_ATTACHMENTS = 3


def safe_file_name(value: str) -> str:
    name = re.split(r"[/\\]", value or "")[-1].strip()
    return (name or "document")[:200]


def extract_document_text(file_name: str, data: bytes) -> tuple[str, str]:
    """Parse supported study documents without executing embedded content."""
    extension = PurePath(file_name.lower()).suffix
    if not data:
        raise HTTPException(status_code=422, detail="文件为空")
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="文件不能超过 5MB")
    try:
        if extension in {".txt", ".md"}:
            text = data.decode("utf-8-sig")
            content_type = "text/markdown" if extension == ".md" else "text/plain"
        elif extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("invalid PDF signature")
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 30:
                raise ValueError("encrypted or over 30 pages")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            content_type = "application/pdf"
        elif extension == ".docx":
            if not zipfile.is_zipfile(io.BytesIO(data)):
                raise ValueError("invalid DOCX archive")
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("DOCX expanded size limit exceeded")
            from docx import Document

            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join([*paragraphs, *cells])
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            raise HTTPException(status_code=422, detail="目前仅支持 TXT、Markdown、PDF 和 DOCX 文档")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="文件无法读取；请检查格式、加密状态或页数") from exc
    text = text.strip()
    if not text:
        raise HTTPException(status_code=422, detail="文件中没有可提取的文字；扫描件暂不支持")
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS] + "\n[仅提取前 20000 字]"
    return content_type, text


def attachment_context(db: Session, user_id: int, attachment_ids: list[int]) -> tuple[str, list[dict[str, object]]]:
    """Read only this user's selected documents and bound model context."""
    if len(attachment_ids) > MAX_ATTACHMENTS or len(set(attachment_ids)) != len(attachment_ids):
        raise HTTPException(status_code=422, detail="每条消息最多选择 3 个不同的附件")
    if not attachment_ids:
        return "", []
    rows = db.query(AssistantAttachment).filter(
        AssistantAttachment.user_id == user_id,
        AssistantAttachment.id.in_(attachment_ids),
    ).all()
    by_id = {row.id: row for row in rows}
    if len(by_id) != len(attachment_ids):
        raise HTTPException(status_code=404, detail="附件不存在或无权访问")
    parts: list[str] = []
    metadata: list[dict[str, object]] = []
    remaining = MAX_PROMPT_CHARS
    for attachment_id in attachment_ids:
        row = by_id[attachment_id]
        metadata.append({"id": row.id, "file_name": row.file_name, "content_type": row.content_type})
        excerpt = row.text_content[:remaining]
        if excerpt:
            parts.append(f"文件：{row.file_name}\n{excerpt}")
            remaining -= len(excerpt)
        if remaining <= 0:
            break
    return "\n\n".join(parts), metadata
