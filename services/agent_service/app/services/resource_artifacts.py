"""Deterministic local PPTX/SVG/WAV production inside the existing assistant worker."""
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import wave
from html import escape
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException

from common.knowledge_paths import KNOWLEDGE_ROOT, PROJECT_ROOT
from common.learning_context import validate_context
from common.models.media import ResourceArtifact
from services.agent_service.app.services.assistant_execution import (
    ensure_execution_active,
    ensure_write_allowed,
)

FORMATS = {"pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
           "animation": ("image/svg+xml", "svg"), "audio": ("audio/wav", "wav"),
           "video": ("video/webm", "webm")}
ARTIFACT_ROOT = KNOWLEDGE_ROOT / "resource_artifacts"
VIDEO_RENDER_VERSION = "math-v2"


def resource_video_path(resource):
    from common.resource_version import resource_version
    return ARTIFACT_ROOT / str(resource.generated_for_user_id or 'public') / f"resource-{resource.id}-{resource_version(resource)}-{VIDEO_RENDER_VERSION}.webm"


def ensure_resource_video(resource):
    """Shared immutable media for every generation entry, independent of a plan."""
    if not resource or resource.type != 'courseware' or resource.generation_status != 'generated':
        raise ValueError('Only completed courseware can produce video')
    target = resource_video_path(resource)
    if not target.is_file():
        payload = render_video(resource.title, resource.content)
        ensure_execution_active()
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(target.name + '.' + uuid4().hex + '.tmp')
        temporary.write_bytes(payload)
        os.replace(temporary, target)
    return {'status': 'ready', 'url': f'/resources/{resource.id}/video',
            'kind': 'caption_video', 'silent': True, 'size': target.stat().st_size}


def render_pptx(title, content):
    from pptx import Presentation
    from pptx.util import Pt
    presentation = Presentation()
    lines = [re.sub(r"^[#*\-\s]+", "", text).strip() for text in content.splitlines() if text.strip()][:200]
    for offset in range(0, max(1, len(lines)), 6):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(title or "学习课件")[:100] if offset == 0 else f"{title or '学习课件'} · {offset//6+1}"
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, text in enumerate(lines[offset:offset+6]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text[:800]
            paragraph.font.size = Pt(20)
    output = io.BytesIO()
    presentation.save(output)
    # Validate real round-trip before allowing persistence.
    assert len(Presentation(io.BytesIO(output.getvalue())).slides) > 0
    return output.getvalue()


def render_animation(title, content):
    steps = [line.lstrip("# -*").strip()[:70] for line in content.splitlines() if line.strip()][:6]
    if not steps:
        raise ValueError("Animation source is empty")
    rows = [f'<text x="35" y="45" font-size="24">{escape(str(title or "知识步骤演示")[:70])}</text>']
    duration = len(steps)*2
    for index, text in enumerate(steps):
        # SVG SMIL only, no script, foreignObject, links or remote assets.
        rows.append(f'<g><rect x="20" y="{75+index*75}" width="860" height="58" rx="10" fill="#dce8ff">'
            f'<animate attributeName="opacity" values="0.25;1;0.25" dur="{duration}s" begin="{index*2}s" repeatCount="indefinite"/></rect>'
            f'<text x="35" y="{112+index*75}" font-size="19">{index+1}. {escape(text)}</text></g>')
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 900 {100+len(steps)*75}" role="img" aria-label="知识步骤动画">'
        '<rect width="100%" height="100%" fill="white"/>' + "".join(rows) + '</svg>').encode()


def render_audio(content):
    if os.name != "nt":
        raise ValueError("BLOCKED: Windows local speech engine is unavailable")
    text = re.sub(r"[#*`>\[\]]", "", content).strip()[:1600]
    if not text:
        raise ValueError("Audio source is empty")
    with tempfile.TemporaryDirectory(prefix="learning-tts-") as directory:
        source, output = Path(directory)/"text.txt", Path(directory)/"speech.wav"
        source.write_text(text, encoding="utf-8")
        result = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-File", str(PROJECT_ROOT/"scripts/render_local_speech.ps1"),
            "-TextPath", str(source), "-OutputPath", str(output)], capture_output=True, timeout=120,
            creationflags=subprocess.CREATE_NO_WINDOW)
        if result.returncode or not output.is_file():
            raise ValueError("BLOCKED: local speech engine did not produce audio")
        with wave.open(str(output)) as audio:
            if audio.getnframes() == 0 or audio.getframerate() <= 0:
                raise ValueError("Local speech produced an invalid WAV")
        return output.read_bytes()


def render_video(title, content):
    """Typeset source Markdown/math offline, then encode a silent local video."""
    import cv2
    import numpy as np
    from PIL import Image, ImageDraw, ImageFont

    # CJK fonts must exist: fail explicitly instead of exporting unreadable boxes.
    fonts = [Path(os.environ.get("WINDIR", "C:/Windows"))/"Fonts/msyh.ttc",
             Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")]
    font_path = next((path for path in fonts if path.is_file()), None)
    if not font_path:
        raise ValueError("视频生成需要可用的中文字体")
    font = ImageFont.truetype(str(font_path), 27)
    heading_font = ImageFont.truetype(str(font_path), 32)
    if not content.strip():
        raise ValueError("视频课件正文为空")
    with tempfile.TemporaryDirectory(prefix="learning-video-") as directory:
        source = Path(directory) / "source.json"
        frames = Path(directory) / "frames"
        source.write_text(json.dumps({"title": title, "content": content}, ensure_ascii=False), encoding="utf-8")
        node = shutil.which("node")
        if not node:
            raise ValueError("视频公式排版需要本机 Node.js 和 Chromium")
        result = subprocess.run([node, str(PROJECT_ROOT / "scripts/render_courseware_frames.mjs"), str(source), str(frames)],
            capture_output=True, timeout=180, creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        if result.returncode:
            raise ValueError("视频公式排版失败：" + result.stderr.decode("utf-8", errors="replace")[-1800:])
        pages = sorted(frames.glob("frame-*.png"))
        if not pages:
            raise ValueError("视频排版没有产生画面")
        target = str(Path(directory)/"lesson.webm")
        writer = cv2.VideoWriter(target, cv2.VideoWriter_fourcc(*"VP80"), 1, (960, 600))
        if not writer.isOpened():
            raise ValueError("本机 VP8 视频编码器不可用")
        extra_frames = 0
        try:
            if "冒泡排序" in str(title) + content[:300]:
                values = [4, 2, 7, 1]
                for end in range(len(values)-1, 0, -1):
                    for left in range(end):
                        ensure_execution_active()
                        swap = values[left] > values[left+1]
                        for after in (False, True):
                            if after and swap:
                                values[left], values[left+1] = values[left+1], values[left]
                            frame = Image.new("RGB", (960, 600), "#f5f8ff")
                            draw = ImageDraw.Draw(frame)
                            draw.text((42, 26), "冒泡排序 · 升序示例 [4, 2, 7, 1]", font=heading_font, fill="#152342")
                            label = ("交换完成" if swap else "保持顺序") if after else f"比较下标 {left} 和 {left+1} 的相邻元素"
                            draw.text((42, 100), label, font=font, fill="#26344d")
                            for index, value in enumerate(values):
                                x = 130+index*180
                                draw.rectangle((x, 470-value*36, x+110, 470), fill="#f0ac40" if index in (left,left+1) else "#5c80d8")
                                draw.text((x+35, 480), str(value), font=heading_font, fill="#152342")
                            draw.text((42, 550), "大于右邻时交换，每轮将最大值移至未排序区末尾 · 无声", font=font, fill="#52688a")
                            encoded = cv2.cvtColor(np.asarray(frame), cv2.COLOR_RGB2BGR)
                            for _ in range(2):
                                writer.write(encoded)
                                extra_frames += 1
            for paragraphs in pages:
                ensure_execution_active()
                with Image.open(paragraphs) as image:
                    frame = image.convert("RGB")
                encoded = cv2.cvtColor(np.asarray(frame), cv2.COLOR_RGB2BGR)
                for _ in range(12):
                    writer.write(encoded)
        finally:
            writer.release()
        capture = cv2.VideoCapture(target)
        try:
            ok, _ = capture.read()
            if not ok or int(capture.get(cv2.CAP_PROP_FRAME_COUNT)) < len(pages)*12+extra_frames:
                raise ValueError("视频文件无法解码或帧数不完整")
        finally:
            capture.release()
        return Path(target).read_bytes()


def produce_artifact(db, task, *, context_id=None, kind=None):
    context_id = context_id or task.request_json.get("learning_context_id")
    kind = kind or task.request_json.get("artifact_kind")
    if kind not in FORMATS:
        raise ValueError("Unsupported artifact type")
    context, _, _, _, _, resource = validate_context(db, task.user_id, context_id)
    if resource is None:
        raise ValueError("Select a resource before creating an artifact")
    existing = db.query(ResourceArtifact).filter_by(context_id=context_id, kind=kind).first()
    if existing:
        verified_path(existing)
        return result_card(existing)
    title, content = resource.title, resource.content
    db.rollback()  # No long SQL transaction during local rendering.
    renderer = {"pptx": render_pptx, "animation": render_animation, "audio": lambda title, content: render_audio(content), "video": render_video}[kind]
    payload = renderer(title, content)
    ensure_write_allowed(db)
    context, _, _, _, _, resource = validate_context(db, task.user_id, context_id)
    existing = db.query(ResourceArtifact).filter_by(context_id=context_id, kind=kind).first()
    if existing:
        return result_card(existing)
    artifact_id = str(uuid4())
    relative = f"{task.user_id}/{artifact_id}.{FORMATS[kind][1]}"
    target = (ARTIFACT_ROOT/relative).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(payload)
    row = ResourceArtifact(id=artifact_id, user_id=task.user_id, resource_id=resource.id, context_id=context_id,
        resource_version=context.resource_version, kind=kind, relative_path=relative, sha256=hashlib.sha256(payload).hexdigest(), size=len(payload))
    db.add(row)
    try:
        db.commit()
    except Exception:
        db.rollback()
        target.unlink(missing_ok=True)
        raise
    return result_card(row)


def result_card(row):
    return {"content": f"已生成并保存{row.kind}学习产物。", "cards": [{"type": "artifact", "artifact_id": row.id,
        "kind": row.kind, "learning_context_id": row.context_id}], "artifact_id": row.id}


def verified_path(row):
    path = (ARTIFACT_ROOT/row.relative_path).resolve()
    if not path.is_relative_to(ARTIFACT_ROOT.resolve()) or not path.is_file():
        raise HTTPException(404, "Artifact file unavailable")
    if hashlib.sha256(path.read_bytes()).hexdigest() != row.sha256:
        raise HTTPException(409, "Artifact integrity check failed")
    return path
