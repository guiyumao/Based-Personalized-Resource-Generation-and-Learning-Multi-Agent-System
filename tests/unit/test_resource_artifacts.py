"""Real local formats, owner checks, revision CAS and Worker write fencing."""
import io
import json

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from common.db.base import Base
from common.db.session import get_db
from common.dependencies import get_current_user
from common.learning_context import create_context
from common.models.learning import KnowledgePoint, LearningPath, Resource, SystemSubject, User
from common.models.media import ResourceArtifact
from common.resource_version import resource_version
from services.agent_service.app.main import app
from services.agent_service.app.services import assistant_execution, assistant_tasks, resource_artifacts


def test_resource_center_mindmap_edit_has_no_plan_requirement_and_checks_owner(media):
    client, factory, _ = media
    data = client.get('/learning/resources/1/mindmap').json()
    body = {'expected_resource_version':data['resource_version'], 'nodes':[{'id':'root','label':'edited TCP','parent_id':None}]}
    assert client.put('/learning/resources/1/mindmap',json=body).status_code == 200
    assert len(client.get('/learning/resources/1/resource-history').json()['items']) == 2
    assert client.put('/learning/resources/1/mindmap',json=body).status_code == 409
    app.dependency_overrides[get_current_user] = lambda: User(id=2,username='other',password_hash='unused')
    assert client.get('/learning/resources/1/mindmap').status_code == 404
    assert client.put('/learning/resources/1/mindmap',json=body).status_code == 404
    assert client.get('/learning/resources/1/resource-history').status_code == 404


def test_resource_video_works_without_learning_plan_and_is_version_scoped(media, monkeypatch):
    from common.db import session
    client, factory, _ = media
    monkeypatch.setattr(session, 'SessionLocal', factory)
    with factory() as db:
        source = db.get(Resource, 1)
        source.type = 'courseware'
        source.content = '# 冒泡排序\n比较相邻元素，大于右邻时交换。'
        db.commit()
        version = resource_version(source)
    assert client.get('/resources/1/video').status_code == 404
    response = client.post('/resources/1/video', json={'client_request_id': 'global-video-test', 'expected_version': version})
    assert response.status_code == 202, response.text
    task = client.get(f"/assistant/tasks/{response.json()['id']}").json()
    assert task['status'] == 'SUCCEEDED', task
    file = client.get('/resources/1/video')
    assert file.status_code == 200 and file.headers['content-type'] == 'video/webm'
    app.dependency_overrides[get_current_user] = lambda: User(id=2, username='other', password_hash='unused')
    assert client.get('/resources/1/video').status_code == 404
    app.dependency_overrides[get_current_user] = lambda: User(id=1, username='owner', password_hash='unused')
    with factory() as db:
        db.get(Resource, 1).content += '\n新的正文'
        db.commit()
    assert client.get('/resources/1/video').status_code == 404
    assert client.post('/resources/1/video', json={'client_request_id': 'stale-video-test', 'expected_version': version}).status_code == 409


@pytest.fixture
def media(tmp_path, monkeypatch):
    engine=create_engine(f"sqlite:///{tmp_path/'media.sqlite'}",connect_args={"check_same_thread":False})
    Base.metadata.create_all(engine)
    factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([User(id=1,username="owner",password_hash="unused"),User(id=2,username="other",password_hash="unused")])
        db.add(SystemSubject(id=1,name="网络",code="network",description="fixture"))
        db.flush()
        db.add(KnowledgePoint(id=1,name="TCP",subject_id=1))
        db.add(LearningPath(id=1,user_id=1,status="active",path_data_json={"subject":"网络","stages":[{"stage_id":"s1","tasks":[{"task_id":"t1","knowledge_point":"TCP"}]}]}))
        db.add(Resource(id=1,title="TCP",type="mind_map",content="# TCP\n确认与重传。",knowledge_point_id=1,generated_for_user_id=1,
            structured_json=json.dumps({"structured":{"nodes":[{"id":"root","label":"TCP","parent_id":None}]}})))
        db.commit()
        context_id=create_context(db,1,client_request_id="media-fixture",plan_id=1,module_id="s1",task_key="t1",resource_id=1,expected_resource_version=resource_version(db.get(Resource,1))).id
    def database():
        with factory() as db:
            yield db
    app.dependency_overrides[get_db]=database
    app.dependency_overrides[get_current_user]=lambda:User(id=1,username="owner",password_hash="unused")
    monkeypatch.setattr(assistant_tasks,"SessionLocal",factory)
    monkeypatch.setattr(assistant_execution,"SessionLocal",factory)
    monkeypatch.setattr(resource_artifacts,"ARTIFACT_ROOT",tmp_path/"artifacts")
    yield TestClient(app), factory, context_id
    app.dependency_overrides.clear()
    engine.dispose()


@pytest.mark.parametrize("kind",["pptx","animation","video"])
def test_actual_file_roundtrip_replay_ownership_and_version(media,kind):
    client,factory,context_id=media
    body={"kind":kind,"client_request_id":f"render-{kind}-fixture"}
    response=client.post(f"/learning/contexts/{context_id}/artifacts",json=body)
    assert response.status_code==202,response.text
    task_id=response.json()["id"]
    result=client.get(f"/assistant/tasks/{task_id}").json()
    assert result["status"]=="SUCCEEDED",result
    artifact_id=result["result"]["artifact_id"]
    content=client.get(f"/learning/artifacts/{artifact_id}/file")
    assert content.status_code==200
    if kind=="pptx":
        from pptx import Presentation
        slides=Presentation(io.BytesIO(content.content)).slides
        assert "TCP" in "".join(shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
    elif kind=="animation":
        assert b"<animate " in content.content and b"<script" not in content.content
    else:
        import cv2
        assert content.headers['content-type'] == 'video/webm'
        assert content.content[:4] == bytes.fromhex('1a45dfa3')
        with factory() as db:
            capture = cv2.VideoCapture(str(resource_artifacts.verified_path(db.get(ResourceArtifact, artifact_id))))
            assert capture.read()[0]
            assert capture.get(cv2.CAP_PROP_FRAME_COUNT) >= 12
            capture.release()
    assert client.post(f"/learning/contexts/{context_id}/artifacts",json=body).json()["id"]==task_id
    app.dependency_overrides[get_current_user]=lambda:User(id=2,username="other",password_hash="unused")
    assert client.get(f"/learning/artifacts/{artifact_id}/file").status_code==404
    app.dependency_overrides[get_current_user]=lambda:User(id=1,username="owner",password_hash="unused")
    with factory() as db:
        db.get(Resource,1).content="new source"
        db.commit()
        assert db.query(ResourceArtifact).count()==1
    assert client.get(f"/learning/artifacts/{artifact_id}/file").status_code==409


def test_mindmap_tree_validation_history_and_old_context_rejection(media):
    client,factory,context_id=media
    with factory() as db:
        version=resource_version(db.get(Resource,1))
    payload={"expected_resource_version":version,"nodes":[{"id":"root","label":"TCP","parent_id":None},{"id":"child","label":"重传","parent_id":"root"}]}
    cycle={**payload,"nodes":[{"id":"root","label":"TCP","parent_id":"child"},{"id":"child","label":"x","parent_id":"root"}]}
    assert client.put(f"/learning/contexts/{context_id}/mindmap",json=cycle).status_code==422
    saved=client.put(f"/learning/contexts/{context_id}/mindmap",json=payload)
    assert saved.status_code==200,saved.text
    assert saved.json()["resource_version"]!=version
    assert client.put(f"/learning/contexts/{context_id}/mindmap",json=payload).status_code==409
    history=client.get(f"/learning/contexts/{context_id}/resource-history")
    assert history.status_code==200 and len(history.json()["items"])==2
    with factory() as db:
        assert "重传" in db.get(Resource,1).content


def test_context_invalidated_during_render_discards_output(media,monkeypatch):
    client,factory,context_id=media
    def change_during_render(*_):
        with factory() as db:
            db.get(Resource,1).publish_status="archived"
            db.commit()
        return b"not allowed to persist"
    monkeypatch.setattr(resource_artifacts,"render_pptx",change_during_render)
    response=client.post(f"/learning/contexts/{context_id}/artifacts",json={"kind":"pptx","client_request_id":"invalidate-render"})
    task=client.get(f"/assistant/tasks/{response.json()['id']}").json()
    assert task["status"]=="FAILED"
    with factory() as db:
        assert db.query(ResourceArtifact).count()==0
