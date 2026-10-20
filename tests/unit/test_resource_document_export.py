"""Real DOCX bytes and private-resource isolation, using an isolated database."""

import io

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from common.db.base import Base
from common.db.session import get_db
from common.models.learning import Resource, User
from common.security.auth import create_access_token
from services.resource_service.app.main import app


def test_word_download_preserves_learning_content_and_enforces_owner():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        owner = User(username="export_owner", password_hash="test", role="student")
        other = User(username="export_other", password_hash="test", role="student")
        db.add_all([owner, other])
        db.flush()
        private = Resource(title="TCP 学习讲义", type="courseware", format="markdown", generated_for_user_id=owner.id,
                           content='# TCP 基础\n**可靠传输**\n公式 $x_i=\\frac{a}{b}$\n```python\nprint("<script>")\n```\n| 协议 | 特点 |\n| --- | --- |\n| TCP | 可靠 |')
        public = Resource(title="公共讲义", type="notes", format="markdown", content="共享课程内容")
        unsupported = Resource(title="结构化资源", type="mind_map", format="json", content="{}")
        db.add_all([private, public, unsupported])
        db.commit()
        private_id, public_id, unsupported_id = private.id, public.id, unsupported.id
        owner_headers = {"Authorization": f"Bearer {create_access_token(str(owner.id))}"}
        other_headers = {"Authorization": f"Bearer {create_access_token(str(other.id))}"}

    def override_db():
        with factory() as db:
            yield db

    previous_override = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override_db
    try:
        client = TestClient(app)
        url = f"/resources/{private_id}/export.docx"
        assert client.get(url).status_code == 401
        assert client.get(url, headers=other_headers).status_code == 404
        response = client.get(url, headers=owner_headers)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert "filename*=UTF-8''" in response.headers["content-disposition"]
        assert response.content.startswith(b"PK")
        document = Document(io.BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "可靠传输" in text
        assert "$x_i=\\frac{a}{b}$" in text
        assert 'print("<script>")' in text
        assert document.tables[0].cell(1, 0).text == "TCP"
        assert client.get(f"/resources/{public_id}/export.docx", headers=other_headers).status_code == 200
        assert client.get(f"/resources/{unsupported_id}/export.docx", headers=owner_headers).status_code == 422
    finally:
        if previous_override is None:
            app.dependency_overrides.pop(get_db, None)
        else:
            app.dependency_overrides[get_db] = previous_override
        engine.dispose()
